#!/usr/bin/env python3
# Provenance: original path 05_AUTOMATION/md2docx.py; original SHA-256 2fa6805420f27cd3cb4d11559db60eb5cbde35f62416ab339308565bfac736a6; classification active (one-offs; OPS 8c1852b); versioned 2026-09-11.
"""Convert GE16 report markdown to formatted .docx (python-docx).

Usage:
  python3 md2docx.py          # EN (default)
  python3 md2docx.py --ms     # Malay (…_MS.md -> …_MS.docx)
  python3 md2docx.py --both   # both languages
"""
import re
import sys
from pathlib import Path


def resolve_repository_root(anchor=None):
    """Find the ANALYTICS repository root independently of cwd."""
    return Path(anchor or __file__).resolve().parents[2]


ROOT = resolve_repository_root()

MS = '--ms' in sys.argv or '--both' in sys.argv
EN = (not MS) or '--both' in sys.argv

JOBS = []
if EN:
    JOBS.append((ROOT / '03_REPORTS/federal/latest/GE16_Malaysia_General_Election_Report.md',
                 ROOT / '03_REPORTS/federal/latest/GE16_Malaysia_General_Election_Report.docx'))
if MS:
    JOBS.append((ROOT / '03_REPORTS/federal/latest/GE16_Malaysia_General_Election_Report_MS.md',
                 ROOT / '03_REPORTS/federal/latest/GE16_Malaysia_General_Election_Report_MS.docx'))


def strip_md(t):
    t = re.sub(r'\*\*(.*?)\*\*', r'\1', t)
    t = re.sub(r'\*(.*?)\*', r'\1', t)
    t = re.sub(r'`(.*?)`', r'\1', t)
    t = re.sub(r'\[(.*?)\]\((.*?)\)', r'\1', t)
    return t


def convert(src, dst):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    def add_md_table(rows):
        ncols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=ncols)
        table.style = 'Light Grid Accent 1'
        for i, row in enumerate(rows):
            for j in range(ncols):
                cell = table.cell(i, j)
                cell.text = row[j] if j < len(row) else ''
                for p in cell.paragraphs:
                    p.style = doc.styles['Normal']
                    for run in p.runs:
                        run.font.size = Pt(9)
        doc.add_paragraph()

    with open(src, encoding='utf-8') as f:
        lines = f.read().split('\n')

    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith('|') and i + 1 < len(lines) and lines[i + 1].strip().startswith('|---'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c or '') for c in cells):
                    rows.append(cells)
                i += 1
            add_md_table(rows)
            continue
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=3)
        elif line.strip() == '---':
            doc.add_paragraph('─' * 60)
        elif line.strip() == '':
            pass
        elif line.strip().startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(strip_md(line.strip()[2:]))
        elif re.match(r'^\d+\.\s', line.strip()):
            p = doc.add_paragraph(style='List Number')
            p.add_run(strip_md(re.sub(r'^\d+\.\s*', '', line.strip())))
        elif line.strip().startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                p = doc.add_paragraph()
                run = p.add_run(lines[i])
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                i += 1
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', strip_md(line))
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part:
                    p.add_run(part)
        i += 1

    doc.save(dst)
    print('DOCX saved:', dst)


def main():
    if '--help' in sys.argv or '-h' in sys.argv:
        print(__doc__)
        return
    for src, dst in JOBS:
        try:
            convert(src, dst)
        except FileNotFoundError:
            print(f'skip (missing): {src}')


if __name__ == '__main__':
    main()
